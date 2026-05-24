"""
build_workbook.py
-----------------
Builds a publication-quality Excel literature-extraction workbook for the topic:
    "Biopolymer Compatibility for Progesterone Encapsulation Using Spray Drying"

Output: Progesterone_SprayDrying_Biopolymer_LitDB.xlsx

Design principles
-----------------
- 10 categorized sheets + README + Quality_Legend
- All sheets use frozen header rows, auto-filters, bold colored headers, wrap-text
- Dropdown data validations on enumerated fields
- Conditional formatting for quality grading and encapsulation efficiency
- "NR" convention for Not Reported; no fabricated study-specific values
- Sheet 8 (comparative biopolymers) and Sheet 10 (field-level insights) are
  pre-populated from canonical review literature; every other sheet is a
  ready-to-fill template seeded with one example row of NR placeholders.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import (
    CellIsRule,
    ColorScaleRule,
    FormulaRule,
)

OUTPUT = "Progesterone_SprayDrying_Biopolymer_LitDB.xlsx"

# ---------------------------------------------------------------------------
# Shared styles
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")          # deep blue
SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")        # mid blue
ZEBRA_FILL = PatternFill("solid", fgColor="F2F2F2")            # light grey
SECTION_FILL = PatternFill("solid", fgColor="305496")          # navy
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")             # pale yellow

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=10, color="1F1F1F")
NOTE_FONT = Font(name="Calibri", size=10, italic=True, color="7F6000")

THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


# ---------------------------------------------------------------------------
# Utility helpers
# ---------------------------------------------------------------------------
def style_header_row(ws, row_idx: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_WRAP
        cell.border = BORDER
    ws.row_dimensions[row_idx].height = 38


def write_headers(ws, headers: list[str], row: int = 1) -> None:
    for i, h in enumerate(headers, start=1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, len(headers))


def set_col_widths(ws, widths: dict[int, int]) -> None:
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def apply_body_style(ws, start_row: int, end_row: int, n_cols: int,
                     center_cols: set[int] | None = None) -> None:
    center_cols = center_cols or set()
    for r in range(start_row, end_row + 1):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER_WRAP if c in center_cols else LEFT_WRAP
            if r % 2 == 0:
                cell.fill = ZEBRA_FILL


def add_autofilter_and_freeze(ws, n_cols: int, n_rows: int) -> None:
    last = f"{get_column_letter(n_cols)}{n_rows}"
    ws.auto_filter.ref = f"A1:{last}"
    ws.freeze_panes = "A2"


def add_dropdown(ws, col_idx: int, options: list[str],
                 first_row: int = 2, last_row: int = 500) -> None:
    formula = '"' + ",".join(options) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=True,
                        errorTitle="Invalid value",
                        error="Please pick from the dropdown list.")
    col_letter = get_column_letter(col_idx)
    dv.add(f"{col_letter}{first_row}:{col_letter}{last_row}")
    ws.add_data_validation(dv)


def fill_nr_row(ws, n_cols: int, row: int, serial: int | None = None,
                title_col: int | None = None) -> None:
    """Seed a row with NR placeholders so the user sees the intended pattern."""
    for c in range(1, n_cols + 1):
        ws.cell(row=row, column=c, value="NR")
    if serial is not None:
        ws.cell(row=row, column=1, value=serial)
    if title_col is not None:
        ws.cell(row=row, column=title_col, value="<Paste article title here>")


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------
def build_readme(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "README"

    ws["A1"] = "Biopolymer Compatibility for Progesterone Encapsulation via Spray Drying"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
    ws.merge_cells("A1:E1")

    ws["A2"] = "Literature extraction & meta-analysis workbook (MSc / PhD / review-writing ready)"
    ws["A2"].font = Font(name="Calibri", size=11, italic=True, color="404040")
    ws.merge_cells("A2:E2")

    rows = [
        ("", ""),
        ("Scope", "Peer-reviewed journal articles, review papers, and indexed publications "
                  "covering progesterone encapsulation, spray drying, biopolymer "
                  "compatibility, controlled release, and reproductive/pharmaceutical delivery."),
        ("Sheet index", ""),
        ("  Sheet 1",  "Master Literature Database  - bibliographic + indexing + quality grade"),
        ("  Sheet 2",  "Study Objectives & Research Gap"),
        ("  Sheet 3",  "Biopolymer Characterization (with dropdowns)"),
        ("  Sheet 4",  "Spray Drying Parameters"),
        ("  Sheet 5",  "Progesterone Formulation Details"),
        ("  Sheet 6",  "Characterization Methods"),
        ("  Sheet 7",  "Release Profile & Biological Performance"),
        ("  Sheet 8",  "Comparative Analysis of Biopolymers (pre-populated reference table)"),
        ("  Sheet 9",  "Key Findings & Conclusions"),
        ("  Sheet 10", "AI-Assisted Research Insights (pre-populated)"),
        ("  Quality_Legend", "Color-coded study quality grading legend"),
        ("", ""),
        ("Conventions", ""),
        ("  NR",            "Not Reported - use this token whenever a study does not state a value."),
        ("  Units",         "Temperatures in degrees C; particle size in micrometres or nanometres "
                           "as reported; always preserve the unit string from the source."),
        ("  Citation key",  "Use first-author surname + year (e.g. Sosnik2014). Re-use this same "
                           "Serial / Citation key across all sheets to keep cross-sheet joins clean."),
        ("  Quality grade", "A = Q1 indexed, IF >= 3, n >= 3, full methodology;  "
                           "B = Q2/Q3 with adequate methods;  "
                           "C = Q4 or limited reporting;  "
                           "F = predatory / non-indexed / suspicious - exclude."),
        ("", ""),
        ("Important rules (per request)", ""),
        ("  *", "Use ONLY peer-reviewed studies."),
        ("  *", "Reject predatory journals; flag suspicious / non-indexed entries with grade 'F'."),
        ("  *", "Do not hallucinate values - leave NR if the article is silent."),
        ("  *", "Maintain citation integrity - keep DOI and PMID exactly as published."),
        ("  *", "Sheet 8 and Sheet 10 carry general reference-grade synthesis from review "
                "literature, NOT study-specific extractions. Always treat them as orientation."),
        ("", ""),
        ("How to use", ""),
        ("  1.", "For every PDF you read, add ONE row to Sheet 1 with a unique serial."),
        ("  2.", "Carry that same serial into Sheets 2-7 and 9, completing the columns "
                "as the article reports."),
        ("  3.", "Use the dropdowns in Sheet 3 (Type, GRAS, Hydrophilicity, etc.) to keep "
                "vocabulary consistent so pivots, filters, and meta-analysis stay clean."),
        ("  4.", "Conditional formatting will automatically colour Encapsulation Efficiency "
                "in Sheet 5 and Quality Grade in Sheet 1 once values are entered."),
        ("  5.", "Use Sheet 8 as your benchmarking baseline; mark which row of Sheet 1 "
                "supports each cell in your own write-up."),
        ("  6.", "Sheet 10 is a living research-gap brief; update it as the database grows."),
        ("", ""),
        ("Citation note for Sheets 8 & 10",
            "Pre-populated reference values draw on widely cited review literature on "
            "spray-drying microencapsulation, biopolymer carriers, and progesterone "
            "delivery, including (representative): Gharsallaoui et al., Food Res Int 2007; "
            "Sosnik & Seremeta, Adv Colloid Interface Sci 2015; Anandharamakrishnan & "
            "Ishwarya, Spray Drying Techniques for Food Ingredient Encapsulation, Wiley "
            "2015; Rathbone et al., Adv Drug Deliv Rev 2002 (intravaginal progesterone); "
            "Estevinho et al., Trends Food Sci Technol 2013. These are starting "
            "anchors - replace / augment with your own indexed citations as you read."),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color="1F4E78")
        ws.cell(row=i, column=2, value=v).alignment = LEFT_WRAP
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)

    ws.column_dimensions["A"].width = 22
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 18
    for r in range(4, 4 + len(rows)):
        ws.row_dimensions[r].height = 30


def build_quality_legend(wb: Workbook) -> None:
    ws = wb.create_sheet("Quality_Legend")
    write_headers(ws, ["Grade", "Criterion", "Color", "Action"])
    legend = [
        ("A", "Q1 indexed (Scopus/WoS), IF >= 3, full methods, n >= 3, statistics reported", "Green",  "Include - high weight"),
        ("B", "Q2 or Q3 indexed, IF 1-3, adequate methods",                                   "Yellow", "Include - moderate weight"),
        ("C", "Q4 indexed but limited methodology / small n / partial data",                  "Orange", "Include with caution"),
        ("F", "Predatory journal, non-indexed, retracted, or suspicious",                     "Red",    "EXCLUDE"),
    ]
    grade_colors = {"A": "C6EFCE", "B": "FFEB9C", "C": "FFD7A8", "F": "FFC7CE"}
    for i, row in enumerate(legend, start=2):
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.font = BODY_FONT
            cell.alignment = LEFT_WRAP
            cell.border = BORDER
        # tint the colour cell
        ws.cell(row=i, column=3).fill = PatternFill("solid", fgColor=grade_colors[row[0]])
    set_col_widths(ws, {1: 8, 2: 70, 3: 14, 4: 30})
    ws.freeze_panes = "A2"


# ----- Sheet 1 -------------------------------------------------------------
def build_master_literature(wb: Workbook) -> None:
    ws = wb.create_sheet("1_Master_Literature")
    headers = [
        "Serial No.", "Article Title", "First Author", "All Authors", "Year",
        "Journal Name", "Publisher", "DOI", "PMID / Scopus ID / WoS ID",
        "Country of Study", "Peer-reviewed (Y/N)", "Indexed Database",
        "Quartile (Q1/Q2/Q3/Q4)", "Impact Factor", "CiteScore", "NAAS Rating",
        "Open Access / Subscription", "PDF File Name", "Direct PDF Hyperlink",
        "Quality Grade (A/B/C/F)", "Notes",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)

    # one example NR row
    fill_nr_row(ws, n_cols, row=2, serial=1, title_col=2)

    # Dropdowns
    add_dropdown(ws, col_idx=11, options=["Yes", "No"])
    add_dropdown(ws, col_idx=12, options=[
        "Scopus", "Web of Science", "PubMed", "DOAJ", "SCIE", "ESCI",
        "PubMed + Scopus", "Scopus + WoS", "PubMed + Scopus + WoS", "Other",
    ])
    add_dropdown(ws, col_idx=13, options=["Q1", "Q2", "Q3", "Q4", "NR"])
    add_dropdown(ws, col_idx=17, options=["Open Access", "Subscription", "Hybrid", "NR"])
    add_dropdown(ws, col_idx=20, options=["A", "B", "C", "F"])

    # Conditional formatting on quality grade column (T = 20)
    grade_rules = {
        "A": ("C6EFCE", "006100"),
        "B": ("FFEB9C", "9C5700"),
        "C": ("FFD7A8", "9C5700"),
        "F": ("FFC7CE", "9C0006"),
    }
    for grade, (fill, font) in grade_rules.items():
        ws.conditional_formatting.add(
            f"T2:T500",
            CellIsRule(operator="equal", formula=[f'"{grade}"'],
                       fill=PatternFill("solid", fgColor=fill),
                       font=Font(color=font, bold=True)),
        )

    # Conditional formatting on Impact Factor (col 14) - color scale
    ws.conditional_formatting.add(
        "N2:N500",
        ColorScaleRule(start_type="num", start_value=0,  start_color="FFC7CE",
                       mid_type="num",   mid_value=3,    mid_color="FFEB9C",
                       end_type="num",   end_value=10,   end_color="C6EFCE"),
    )

    set_col_widths(ws, {
        1: 9, 2: 55, 3: 18, 4: 35, 5: 7, 6: 28, 7: 22, 8: 28, 9: 22,
        10: 16, 11: 12, 12: 22, 13: 10, 14: 10, 15: 10, 16: 10,
        17: 18, 18: 24, 19: 28, 20: 10, 21: 30,
    })
    apply_body_style(ws, 2, 30, n_cols, center_cols={1, 5, 11, 13, 14, 15, 16, 17, 20})
    add_autofilter_and_freeze(ws, n_cols, 30)


# ----- Sheet 2 -------------------------------------------------------------
def build_objectives(wb: Workbook) -> None:
    ws = wb.create_sheet("2_Objectives_Gap")
    headers = [
        "Serial No.", "Article Title (short)", "Aim / Objectives", "Hypothesis",
        "Research Gap Identified", "Why Spray Drying Was Selected",
        "Novelty of Study", "Scientific Importance", "Industrial Relevance",
        "Veterinary / Reproductive Application",
        "Future Recommendations by Authors",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)
    fill_nr_row(ws, n_cols, row=2, serial=1, title_col=2)

    set_col_widths(ws, {
        1: 9, 2: 35, 3: 42, 4: 32, 5: 38, 6: 32, 7: 32,
        8: 32, 9: 32, 10: 36, 11: 38,
    })
    apply_body_style(ws, 2, 30, n_cols, center_cols={1})
    add_autofilter_and_freeze(ws, n_cols, 30)


# ----- Sheet 3 -------------------------------------------------------------
def build_biopolymer(wb: Workbook) -> None:
    ws = wb.create_sheet("3_Biopolymer_Char")
    headers = [
        "Serial No.", "Biopolymer Name",
        "Type (Natural / Synthetic / Semi-synthetic)",
        "Source of Biopolymer", "Grade (Food / Pharma / Reagent)",
        "Biocompatibility", "Biodegradability", "GRAS Status",
        "Mucoadhesive Property", "Hydrophilic / Hydrophobic",
        "Molecular Weight (Da or kDa)", "Glass Transition Temp (deg C)",
        "Solubility", "Viscosity (cP / mPa.s)",
        "Polymer Concentration (% w/v)",
        "Crosslinking Agent", "Surface Modifier",
        "Emulsifier / Surfactant", "Notes / Reference",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)
    fill_nr_row(ws, n_cols, row=2, serial=1)

    add_dropdown(ws, 3, ["Natural", "Synthetic", "Semi-synthetic", "Blend", "NR"])
    add_dropdown(ws, 5, ["Food-grade", "Pharmaceutical-grade", "Reagent-grade",
                         "Cosmetic-grade", "NR"])
    add_dropdown(ws, 6, ["High", "Moderate", "Low", "NR"])
    add_dropdown(ws, 7, ["High", "Moderate", "Low", "Non-biodegradable", "NR"])
    add_dropdown(ws, 8, ["GRAS", "Not GRAS", "NR"])
    add_dropdown(ws, 9, ["Yes", "No", "NR"])
    add_dropdown(ws, 10, ["Hydrophilic", "Hydrophobic", "Amphiphilic", "NR"])

    set_col_widths(ws, {
        1: 9, 2: 22, 3: 22, 4: 22, 5: 18, 6: 14, 7: 14, 8: 12, 9: 14,
        10: 18, 11: 18, 12: 18, 13: 22, 14: 16, 15: 16, 16: 18,
        17: 18, 18: 22, 19: 32,
    })
    apply_body_style(ws, 2, 30, n_cols, center_cols={1, 3, 5, 6, 7, 8, 9, 10})
    add_autofilter_and_freeze(ws, n_cols, 30)


# ----- Sheet 4 -------------------------------------------------------------
def build_spray_drying(wb: Workbook) -> None:
    ws = wb.create_sheet("4_Spray_Drying_Params")
    headers = [
        "Serial No.", "Spray Dryer Model", "Manufacturer", "Drying Technique",
        "Inlet Temp (deg C)", "Outlet Temp (deg C)",
        "Feed Flow Rate (mL/min or g/min)",
        "Atomization Pressure (bar)", "Nozzle Type",
        "Drying Air Flow (m3/h)", "Aspirator Setting (% / m3/h)",
        "Feed Solid Content (% w/v)", "Solvent System",
        "Emulsion Type", "Drying Yield (%)", "Batch Size",
        "Encapsulation Process Details", "Pre-treatment Method",
        "Notes / Reference",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)
    fill_nr_row(ws, n_cols, row=2, serial=1)

    add_dropdown(ws, 4, [
        "Co-current spray drying", "Counter-current spray drying",
        "Mixed-flow spray drying", "Nano spray drying",
        "Ultrasonic spray drying", "NR",
    ])
    add_dropdown(ws, 9, [
        "Two-fluid nozzle", "Pressure nozzle", "Rotary atomizer",
        "Ultrasonic nozzle", "Three-fluid nozzle", "Vibrating mesh", "NR",
    ])
    add_dropdown(ws, 14, [
        "O/W", "W/O", "W/O/W", "O/W/O", "Single solution", "Suspension", "NR",
    ])

    set_col_widths(ws, {
        1: 9, 2: 22, 3: 20, 4: 24, 5: 14, 6: 14, 7: 18, 8: 16,
        9: 22, 10: 16, 11: 18, 12: 18, 13: 22, 14: 18, 15: 14,
        16: 14, 17: 38, 18: 24, 19: 30,
    })
    apply_body_style(ws, 2, 30, n_cols, center_cols={1, 5, 6, 7, 8, 11, 12, 15, 16})
    add_autofilter_and_freeze(ws, n_cols, 30)


# ----- Sheet 5 -------------------------------------------------------------
def build_progesterone(wb: Workbook) -> None:
    ws = wb.create_sheet("5_Progesterone_Formulation")
    headers = [
        "Serial No.", "Progesterone Source", "Purity (%)",
        "Hormone Concentration", "Drug : Polymer Ratio",
        "Encapsulation Technique", "Encapsulation Efficiency (%)",
        "Loading Capacity (%)", "Particle Size (um or nm)",
        "Zeta Potential (mV)", "Surface Morphology", "PDI",
        "Moisture Content (%)", "Flowability", "Hygroscopicity",
        "Stability Data", "Thermal Stability",
        "Oxidative Stability", "Storage Conditions", "Shelf Life",
        "Notes / Reference",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)
    fill_nr_row(ws, n_cols, row=2, serial=1)

    add_dropdown(ws, 6, [
        "Spray drying (single-stage)", "Spray drying + emulsion",
        "Spray drying + ionic gelation", "Spray drying + coacervation",
        "Spray congealing", "NR",
    ])
    add_dropdown(ws, 11, [
        "Spherical, smooth", "Spherical, dimpled", "Spherical, wrinkled",
        "Irregular", "Hollow", "Aggregated", "NR",
    ])
    add_dropdown(ws, 14, ["Free-flowing", "Cohesive", "Very cohesive", "NR"])

    # Conditional formatting on Encapsulation Efficiency (col 7)
    ws.conditional_formatting.add(
        "G2:G500",
        ColorScaleRule(start_type="num", start_value=0,  start_color="FFC7CE",
                       mid_type="num",   mid_value=70,   mid_color="FFEB9C",
                       end_type="num",   end_value=100,  end_color="C6EFCE"),
    )

    set_col_widths(ws, {
        1: 9, 2: 22, 3: 10, 4: 18, 5: 16, 6: 28, 7: 18, 8: 14,
        9: 18, 10: 16, 11: 22, 12: 10, 13: 14, 14: 14, 15: 14,
        16: 28, 17: 18, 18: 18, 19: 24, 20: 14, 21: 30,
    })
    apply_body_style(ws, 2, 30, n_cols,
                     center_cols={1, 3, 5, 7, 8, 9, 10, 12, 13, 14, 15, 17, 18, 20})
    add_autofilter_and_freeze(ws, n_cols, 30)


# ----- Sheet 6 -------------------------------------------------------------
def build_characterization(wb: Workbook) -> None:
    ws = wb.create_sheet("6_Characterization_Methods")
    headers = [
        "Serial No.", "SEM Analysis", "TEM Analysis", "FTIR", "DSC", "XRD",
        "HPLC", "UV Spectroscopy", "GC-MS", "Rheology",
        "Release Kinetics Method", "Statistical Software",
        "Statistical Tests Used", "Replication Number (n)",
        "Significance Level (alpha)", "Notes / Reference",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)
    fill_nr_row(ws, n_cols, row=2, serial=1)

    yes_no = ["Yes", "No", "NR"]
    for col in (2, 3, 4, 5, 6, 7, 8, 9, 10):
        add_dropdown(ws, col, yes_no)

    set_col_widths(ws, {
        1: 9, 2: 12, 3: 12, 4: 10, 5: 10, 6: 10, 7: 10, 8: 16,
        9: 10, 10: 12, 11: 26, 12: 22, 13: 28, 14: 14, 15: 16, 16: 30,
    })
    apply_body_style(ws, 2, 30, n_cols, center_cols=set(range(2, 11)) | {1, 14, 15})
    add_autofilter_and_freeze(ws, n_cols, 30)


# ----- Sheet 7 -------------------------------------------------------------
def build_release_profile(wb: Workbook) -> None:
    ws = wb.create_sheet("7_Release_Bioperformance")
    headers = [
        "Serial No.", "In vitro Release Duration", "Release Pattern",
        "Burst Release Observed (Y/N)", "Sustained Release (Y/N)",
        "Controlled Release (Y/N)", "Release Kinetics Model",
        "Bioavailability", "Hormone Stability",
        "Plasma Progesterone Data", "In vivo Trial (Y/N)",
        "Animal Species", "Route of Administration",
        "Reproductive Response", "Fertility Outcomes",
        "Estrus Synchronization Effects", "Pregnancy Rate (%)",
        "Toxicity Findings", "Cytotoxicity Results",
        "Notes / Reference",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)
    fill_nr_row(ws, n_cols, row=2, serial=1)

    yn = ["Yes", "No", "NR"]
    add_dropdown(ws, 4, yn)
    add_dropdown(ws, 5, yn)
    add_dropdown(ws, 6, yn)
    add_dropdown(ws, 11, yn)
    add_dropdown(ws, 7, [
        "Zero-order", "First-order", "Higuchi", "Korsmeyer-Peppas",
        "Hixson-Crowell", "Weibull", "Baker-Lonsdale", "NR",
    ])
    add_dropdown(ws, 13, [
        "Oral", "Intravaginal", "Intramuscular", "Subcutaneous",
        "Transdermal", "Intrauterine", "Intranasal", "NR",
    ])

    set_col_widths(ws, {
        1: 9, 2: 18, 3: 22, 4: 14, 5: 14, 6: 14, 7: 22, 8: 18,
        9: 18, 10: 24, 11: 12, 12: 18, 13: 20, 14: 24, 15: 22,
        16: 26, 17: 16, 18: 22, 19: 22, 20: 30,
    })
    apply_body_style(ws, 2, 30, n_cols,
                     center_cols={1, 4, 5, 6, 11, 17})
    add_autofilter_and_freeze(ws, n_cols, 30)


# ----- Sheet 8 -------------------------------------------------------------
def build_comparative_biopolymers(wb: Workbook) -> None:
    """
    Reference-grade comparative table of biopolymers used in spray-dried
    progesterone (and related lipophilic active) microencapsulation.

    Values are qualitative orientation derived from canonical review
    literature (see README citation note). They are NOT extracted from
    a single primary study and must be replaced / augmented with
    study-specific numbers from Sheets 1-7 as your database grows.
    """
    ws = wb.create_sheet("8_Comparative_Biopolymers")

    headers = [
        "Biopolymer", "Class",
        "Encapsulation Efficiency (typical range)",
        "Stability of Spray-Dried Powder",
        "Cost-effectiveness",
        "Release Behaviour (typical)",
        "Biocompatibility",
        "Industrial Scalability for Spray Drying",
        "Ease of Spray Drying",
        "Reproductive / Veterinary Application Suitability",
        "Notes & Caveats",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)

    # qualitative reference rows -------------------------------------------------
    # (Encapsulation efficiency ranges are typical published bands for
    # spray-dried lipophilic / hormone-like actives; all marked "approx."
    # to flag they are reference-band, not point estimates.)
    rows = [
        ("Chitosan", "Natural cationic polysaccharide",
         "approx. 60-90% (lipophilic actives)",
         "Good in dry state, hygroscopic if salt form",
         "Moderate",
         "pH-responsive, mucoadhesive, often sustained",
         "High; biodegradable",
         "Good (low solubility above pH 6 needs acid feed)",
         "Moderate - viscosity rises sharply with concentration",
         "High - mucoadhesion suits intravaginal / intrauterine progesterone delivery",
         "Often blended with alginate or maltodextrin to improve drying yield."),
        ("Alginate", "Natural anionic polysaccharide",
         "approx. 55-85%",
         "Sensitive to humidity unless crosslinked (Ca2+)",
         "Low - moderate",
         "Diffusion + erosion; fast in alkaline media",
         "High; GRAS",
         "Good - widely used at commercial scale",
         "Moderate - can clog nozzles at high viscosity",
         "Good - used in vaginal inserts and bolus matrices",
         "Ionic crosslinking (Ca2+) before/after spray drying improves retention."),
        ("Gelatin", "Natural protein (collagen-derived)",
         "approx. 65-90%",
         "Good if low residual moisture; thermoplastic risk",
         "Moderate",
         "Rapid in aqueous media; sustained when crosslinked",
         "High; biodegradable",
         "Good",
         "Easy - films well around oil droplets",
         "Moderate - fast release limits long-acting progesterone use",
         "Frequently combined with gum arabic / maltodextrin to balance Tg."),
        ("Maltodextrin", "Semi-synthetic carbohydrate (starch hydrolysate)",
         "approx. 50-85% (poor alone for lipophilic actives)",
         "Excellent - high Tg powders",
         "Low (very economical)",
         "Fast dissolution; minimal sustained release alone",
         "High; GRAS",
         "Excellent - industry default carrier",
         "Excellent - low viscosity, high solid loading",
         "Carrier role - rarely the sole wall for progesterone",
         "Use as bulking carrier; combine with gum arabic or whey for emulsifying capacity."),
        ("Gum Arabic (Acacia)", "Natural polysaccharide-protein complex",
         "approx. 70-93% (excellent for lipophiles)",
         "Excellent",
         "Moderate-high (price volatile)",
         "Moderate; matrix-controlled",
         "High; GRAS",
         "Excellent",
         "Excellent - amphiphilic, low viscosity at high solids",
         "Good - benchmark wall for hydrophobic actives incl. progesterone",
         "Often the gold-standard reference wall material in spray-drying studies."),
        ("Whey Protein (WPI / WPC)", "Natural protein",
         "approx. 70-95%",
         "Good; moisture-sensitive at low Tg",
         "Moderate",
         "Stomach-soluble; pepsin-cleavable matrix",
         "High; GRAS",
         "Excellent",
         "Good - strong emulsification, foaming at high shear",
         "Good - relevant to oral / feed-additive progesterone",
         "Heat treatment of feed (denaturation) can boost EE."),
        ("PLGA", "Synthetic biodegradable polyester",
         "approx. 60-92%",
         "Excellent; very low moisture uptake",
         "High (expensive)",
         "Tunable sustained release (weeks-months)",
         "High; FDA-approved for parenterals",
         "Moderate - solvent recovery + GMP costly",
         "Moderate - requires organic solvents (DCM, ACN)",
         "Excellent for injectable long-acting progesterone microspheres",
         "The preferred polymer for parenteral controlled-release hormone depots."),
        ("Pectin", "Natural anionic polysaccharide",
         "approx. 55-82%",
         "Good if low moisture",
         "Low-moderate",
         "Colon-targeted; degraded by colonic microbiota",
         "High; GRAS",
         "Good",
         "Moderate - high viscosity at >5% w/v",
         "Niche - oral colon-targeted delivery",
         "Useful for oral progesterone with colonic absorption strategies."),
        ("Starch (native & modified)", "Natural / semi-synthetic polysaccharide",
         "approx. 50-80%",
         "Good (high Tg in modified forms)",
         "Low (very economical)",
         "Slow swelling; enzyme-dependent",
         "High; GRAS",
         "Excellent",
         "Modified starches (OSA-starch) excellent; native starch poor emulsifier",
         "Moderate - mostly carrier role for progesterone",
         "OSA-modified starch (e.g., HiCap, Capsul) is a strong wall for lipophilic hormones."),
        ("Cellulose derivatives (HPMC, EC, CMC)", "Semi-synthetic polysaccharide",
         "approx. 60-90%",
         "Excellent",
         "Moderate",
         "Tunable - HPMC swells, EC sustains",
         "High",
         "Good (pharma-grade widely available)",
         "Good - HPMC; EC needs organic solvents",
         "Good for oral controlled-release progesterone tablets after spray drying",
         "Combine HPMC + EC for biphasic burst-then-sustained progesterone release."),
        ("Cyclodextrins (alpha, beta, HP-beta-CD)", "Semi-synthetic cyclic oligosaccharide",
         "approx. 70-95% as inclusion complex",
         "Excellent",
         "Moderate-high (HP-beta-CD pricier)",
         "Rapid dissolution; solubility enhancement dominates",
         "High; HP-beta-CD pharma-grade",
         "Good",
         "Excellent - low viscosity solutions",
         "Excellent for solubility/bioavailability boost of progesterone (hydrophobic)",
         "Forms 1:1 inclusion complex with progesterone steroid skeleton."),
        ("Sodium Caseinate", "Natural protein",
         "approx. 70-93%",
         "Good",
         "Moderate",
         "Gastric-soluble; sustained in dairy matrices",
         "High; GRAS",
         "Excellent",
         "Excellent emulsifier at low concentrations",
         "Good - relevant to feed/oral progesterone in ruminants",
         "Strong synergy with maltodextrin as combined wall."),
        ("Carrageenan", "Natural sulphated polysaccharide",
         "approx. 55-85%",
         "Good if low moisture",
         "Moderate",
         "Thermoreversible gel; pH-modulated",
         "High; GRAS (kappa, iota)",
         "Moderate",
         "Moderate - high viscosity above 1% w/v",
         "Niche - vaginal gel-forming progesterone systems",
         "Often blended with chitosan for polyelectrolyte complex."),
        ("Pullulan", "Natural exopolysaccharide (Aureobasidium)",
         "approx. 65-92%",
         "Excellent - very low oxygen permeability",
         "High (expensive)",
         "Slow erosion; oxygen-barrier matrix",
         "High; GRAS",
         "Moderate - cost limits scale",
         "Excellent - films around droplets cleanly",
         "Promising for oxygen-sensitive hormone delivery",
         "Best for protecting progesterone from oxidative degradation."),
        ("Xanthan Gum", "Natural microbial polysaccharide",
         "approx. 50-78%",
         "Good",
         "Low-moderate",
         "Diffusion-controlled; high viscosity matrix",
         "High; GRAS",
         "Moderate - viscosity limits feed solids",
         "Difficult - very viscous at low concentration",
         "Niche - oral mucoadhesive progesterone",
         "Usually a co-polymer (with maltodextrin) rather than sole wall."),
        ("Other (Inulin, Zein, Soy protein, Arabic-WPC blends, etc.)",
         "Mixed",
         "approx. 60-90% (formulation-dependent)",
         "Variable",
         "Variable",
         "Variable",
         "High (mostly food-grade)",
         "Variable",
         "Variable - blends often outperform single polymers",
         "Strong candidate space for novel veterinary progesterone systems",
         "Document blends in Sheet 3 explicitly (each polymer = one row)."),
    ]

    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)

    set_col_widths(ws, {
        1: 28, 2: 30, 3: 30, 4: 30, 5: 18, 6: 32,
        7: 22, 8: 32, 9: 36, 10: 42, 11: 50,
    })
    apply_body_style(ws, 2, 1 + len(rows), n_cols)

    # Note row at the bottom
    note_row = 1 + len(rows) + 2
    note = ("Reference orientation only. Values synthesised from canonical "
            "spray-drying / encapsulation review literature - replace with "
            "study-specific numbers as you populate Sheets 1-7. See README "
            "citation note. Do not cite this sheet directly.")
    ws.cell(row=note_row, column=1, value="NOTE").font = NOTE_FONT
    ws.cell(row=note_row, column=2, value=note).font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=2,
                   end_row=note_row, end_column=n_cols)
    ws.cell(row=note_row, column=1).fill = NOTE_FILL
    ws.cell(row=note_row, column=2).fill = NOTE_FILL
    ws.cell(row=note_row, column=2).alignment = LEFT_WRAP
    ws.row_dimensions[note_row].height = 60

    add_autofilter_and_freeze(ws, n_cols, 1 + len(rows))


# ----- Sheet 9 -------------------------------------------------------------
def build_findings(wb: Workbook) -> None:
    ws = wb.create_sheet("9_Findings_Conclusions")
    headers = [
        "Serial No.", "Major Findings", "Best Performing Biopolymer",
        "Most Efficient Spray Drying Condition", "Best Release Profile",
        "Limitations", "Author Conclusions", "Practical Recommendations",
        "Scope for Future Research",
    ]
    write_headers(ws, headers)
    n_cols = len(headers)
    fill_nr_row(ws, n_cols, row=2, serial=1)

    set_col_widths(ws, {
        1: 9, 2: 50, 3: 26, 4: 38, 5: 30, 6: 38, 7: 42, 8: 38, 9: 38,
    })
    apply_body_style(ws, 2, 30, n_cols, center_cols={1})
    add_autofilter_and_freeze(ws, n_cols, 30)


# ----- Sheet 10 ------------------------------------------------------------
def build_ai_insights(wb: Workbook) -> None:
    ws = wb.create_sheet("10_AI_Research_Insights")

    title = ("AI-Assisted Research Insights - Progesterone microencapsulation "
             "by spray drying (field-level synthesis)")
    ws.cell(row=1, column=1, value=title).font = SECTION_FONT
    ws.cell(row=1, column=1).fill = SECTION_FILL
    ws.cell(row=1, column=1).alignment = LEFT_WRAP
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 28

    headers = ["Theme", "Insight", "Evidence Anchor",
               "Suggested Action / Research Direction"]
    write_headers(ws, headers, row=2)
    n_cols = len(headers)

    # All insights are field-level orientation; not study-specific extractions.
    insights = [
        ("Most frequently used wall polymers",
         "Maltodextrin, gum arabic, whey protein, sodium caseinate, chitosan, "
         "and modified starches dominate the spray-drying literature; PLGA "
         "dominates parenteral long-acting progesterone microspheres.",
         "Cross-tabulate Sheet 3 by polymer name once populated.",
         "Quantify polymer frequency with a pivot on Sheet 3."),

        ("Highest reported encapsulation efficiencies",
         "Polymer blends (gum arabic + maltodextrin; WPI + maltodextrin; "
         "OSA-starch + maltodextrin) routinely outperform single-polymer walls "
         "for lipophilic steroids.",
         "Triangulate from Sheet 5 EE column once populated.",
         "Plan a 2x2 or 3-component mixture-design study comparing best blends."),

        ("Spray-drying parameter consensus",
         "Inlet 140-180 deg C, outlet 70-95 deg C, feed solids 10-30% w/v, "
         "two-fluid nozzle, atomization 2-6 bar are the central tendency for "
         "progesterone-grade systems with thermolabile actives.",
         "Pivot Sheet 4 by inlet/outlet temp to confirm in your own corpus.",
         "Run a CCD/RSM around this central operating window."),

        ("Recurrent methodological weaknesses",
         "Many studies omit (i) detailed in-vivo plasma kinetics, (ii) batch-to-"
         "batch reproducibility data (n<3), (iii) long-term stability beyond "
         "3 months, (iv) species-relevant in-vivo validation in ruminants.",
         "Use Sheet 6 (replication n) and Sheet 7 (in-vivo Y/N) as filters.",
         "Pre-register replication n>=3 and >=12-month accelerated stability."),

        ("Underexplored research gaps",
         "(a) Veterinary intravaginal CIDR-equivalent spray-dried inserts; "
         "(b) ruminant-specific oral bypass coatings; (c) cyclodextrin-"
         "progesterone inclusion + spray-drying hybrid systems; "
         "(d) sustainable plant-protein walls (pea, faba) for hormone delivery.",
         "Filter Sheet 7 for empty 'Animal Species' / 'Reproductive Response'.",
         "Each gap is a candidate MSc/PhD chapter."),

        ("Emerging trends",
         "Nano spray drying (Buchi B-90), three-fluid nozzles for core-shell "
         "particles, and protein-polysaccharide Maillard-conjugate walls are "
         "the strongest recent vectors.",
         "Filter Sheet 4 'Drying Technique' for nano / three-fluid entries.",
         "Benchmark conventional B-290 vs B-90 for progesterone EE and PSD."),

        ("Suggested novel formulation combinations",
         "(1) HP-beta-CD/progesterone inclusion complex + WPI-maltodextrin "
         "spray-dried wall; (2) Chitosan-alginate polyelectrolyte complex with "
         "OSA-starch carrier; (3) Pullulan + zein composite for oxidative "
         "protection; (4) PLGA microspheres with mannitol carrier (parenteral).",
         "These intersect Sheet 8 columns 'EE' and 'Reproductive suitability'.",
         "Each combination = one DoE-driven thesis aim."),

        ("Potential thesis topics",
         "(i) Comparative spray-drying compatibility of biopolymers for "
         "long-acting progesterone in dairy buffalo; "
         "(ii) Cyclodextrin-assisted spray-dried progesterone for intravaginal "
         "estrus synchronization in ewes; "
         "(iii) RSM optimization of WPI-maltodextrin-progesterone microspheres; "
         "(iv) In-vivo plasma kinetics of spray-dried progesterone in repeat-"
         "breeder cows.",
         "Maps directly to Sheet 7 'Reproductive Response' gaps.",
         "Pick one and define the DoE in Sheet 2 'Aim/Objectives'."),

        ("Potential publication ideas",
         "(a) Systematic review: 'Biopolymer compatibility for spray-dried "
         "progesterone - a 2010-2025 meta-analysis'; "
         "(b) Original article: 'OSA-starch vs gum arabic walls for spray-"
         "dried progesterone microparticles'; "
         "(c) Short communication: 'Effect of inlet temperature on residual "
         "progesterone potency post spray drying'; "
         "(d) Veterinary paper: 'Estrus response in repeat-breeder cows to "
         "spray-dried progesterone insert'.",
         "Each paper draws cleanly from one or two filters across Sheets 1-9.",
         "Sheet 1 + Sheet 8 + Sheet 10 already give you the review skeleton."),

        ("Meta-analysis readiness",
         "Once Sheet 5 has >=20 populated rows with EE, particle size and PDI, "
         "a forest plot of EE by polymer class becomes feasible.",
         "Trigger when count(Sheet 5 EE not NR) >= 20.",
         "Use random-effects model (REML); export Sheet 5 to R/metafor."),
    ]

    for r, row in enumerate(insights, start=3):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)

    apply_body_style(ws, 3, 2 + len(insights), n_cols)

    # final caveat note
    note_row = 3 + len(insights) + 1
    ws.cell(row=note_row, column=1, value="NOTE").font = NOTE_FONT
    ws.cell(row=note_row, column=1).fill = NOTE_FILL
    note = ("Field-level orientation generated to seed your research-gap analysis. "
            "Replace each insight with study-specific evidence (Serial No.) as "
            "Sheets 1-7 are populated. Do not cite this sheet directly.")
    ws.cell(row=note_row, column=2, value=note).font = NOTE_FONT
    ws.cell(row=note_row, column=2).fill = NOTE_FILL
    ws.cell(row=note_row, column=2).alignment = LEFT_WRAP
    ws.merge_cells(start_row=note_row, start_column=2,
                   end_row=note_row, end_column=n_cols)
    ws.row_dimensions[note_row].height = 50

    set_col_widths(ws, {1: 32, 2: 70, 3: 38, 4: 50})
    add_autofilter_and_freeze(ws, n_cols, 2 + len(insights))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> None:
    wb = Workbook()
    build_readme(wb)
    build_master_literature(wb)
    build_objectives(wb)
    build_biopolymer(wb)
    build_spray_drying(wb)
    build_progesterone(wb)
    build_characterization(wb)
    build_release_profile(wb)
    build_comparative_biopolymers(wb)
    build_findings(wb)
    build_ai_insights(wb)
    build_quality_legend(wb)

    # Reorder so README is first, then numbered sheets in order, then legend
    desired = [
        "README",
        "1_Master_Literature",
        "2_Objectives_Gap",
        "3_Biopolymer_Char",
        "4_Spray_Drying_Params",
        "5_Progesterone_Formulation",
        "6_Characterization_Methods",
        "7_Release_Bioperformance",
        "8_Comparative_Biopolymers",
        "9_Findings_Conclusions",
        "10_AI_Research_Insights",
        "Quality_Legend",
    ]
    wb._sheets = [wb[name] for name in desired]

    wb.save(OUTPUT)
    print(f"OK  Workbook written: {OUTPUT}")
    print(f"    Sheets: {len(wb.sheetnames)}")
    for s in wb.sheetnames:
        print(f"      - {s}")


if __name__ == "__main__":
    main()
