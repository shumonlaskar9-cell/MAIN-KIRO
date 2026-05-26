"""
Generate the NDRI "FORMAT FOR APPROVAL OF TITLE AND ADVISORY COMMITTEE"
.docx file for the PhD research:
  "Biopolymer Encapsulated Micronised Progesterone Delivery System for
   the Management of Anestrus and Repeat Breeding Syndrome in Cattle"

Major Advisor: Dr. Jeyakumar Sakthivel
Major subject:  ANIMAL REPRODUCTION, GYNAECOLOGY AND OBSTETRICS (ARGO)
Allied:         Animal Biochemistry, Dairy Engineering
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------- helpers ----------

def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)
    tc_pr.append(tcBorders)


def add_run(para, text, bold=False, italic=False, size=12,
            font='Times New Roman', color=None):
    r = para.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def justified_para(doc, line_spacing=1.5, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    return p


def labelled_field(doc, label, value, bold_value=False):
    """Field-label paragraph used in NDRI student-details block."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    add_run(p, f'{label} ', bold=True, size=12)
    add_run(p, value, bold=bold_value, size=12)
    return p


# ---------- document ----------

doc = Document()

# default style
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

# A4, 1-inch margins
for section in doc.sections:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


# ===== HEADER =====
h1 = doc.add_paragraph()
h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(h1, 'NATIONAL DAIRY RESEARCH INSTITUTE (DEEMED UNIVERSITY)',
        bold=True, size=14)
h2 = doc.add_paragraph()
h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(h2, 'KARNAL-132001 (HARYANA)', bold=True, size=14)

doc.add_paragraph()  # blank

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(t, 'FORMAT FOR APPROVAL OF TITLE AND ADVISORY COMMITTEE',
        bold=True, size=13)

doc.add_paragraph()


# ===== STUDENT DETAILS =====
labelled_field(doc, 'Name of the student:', '[STUDENT NAME — to be filled]')
labelled_field(doc, 'Registration No:', '[REG NO — to be filled]')
labelled_field(doc, 'Major subject:',
               'ANIMAL REPRODUCTION, GYNAECOLOGY AND OBSTETRICS (ARGO)')
labelled_field(doc, 'Allied Subject:',
               'Animal Biochemistry, Dairy Engineering.')
labelled_field(doc, 'Major advisor:', 'Dr. Jeyakumar Sakthivel')


# ===== PROJECT OF THE FACULTY =====
pf = doc.add_paragraph()
pf.paragraph_format.space_after = Pt(2)
add_run(pf, 'Project of the Faculty (IRC/Externally funded project):',
        bold=True, size=12)

projects = [
    ('Externally funded project (NASF-CCF-ICAR): Central Corpus Fund - '
     'Indian Council of Agricultural Research (CCF-ICAR), titled "Development '
     'of a novel spray dry based progesterone nanoencapsulation for estrus '
     'synchronization in dairy cattle" (Co-PI)'),
    ('IRC Project no. B-63, titled "Development and Characterization of '
     'Progesterone Loaded Nanofibre for Controlled Breeding in Dairy Cattle" '
     '(Co-PI)'),
    ('IRC Project no. A-97, titled "Validation of progesterone loaded '
     'nanofiber based polyurethane sponge system for estrus synchronization '
     'in dairy cattle" (Co-PI)'),
]
for i, txt in enumerate(projects, start=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p, f'{i}. ', bold=True)
    add_run(p, txt)


# ===== TITLE =====
ttl = doc.add_paragraph()
ttl.paragraph_format.space_before = Pt(6)
add_run(ttl, 'Tentative Title of the Research Project of the student:',
        bold=True, size=12)

ttl_v = justified_para(doc, line_spacing=1.5, space_after=6)
add_run(ttl_v,
        '"Biopolymer Encapsulated Micronised Progesterone Delivery System '
        'for the Management of Anestrus and Repeat Breeding Syndrome in '
        'Cattle"', bold=True, italic=True, size=12)


# ===== OBJECTIVES =====
obj_h = doc.add_paragraph()
obj_h.paragraph_format.space_before = Pt(6)
add_run(obj_h, 'Objectives:', bold=True, size=12)

objectives = [
    ('Optimization of select polymer (sodium alginate) for progesterone '
     'nano/micro-encapsulation by spray drying method.'),
    ('Evaluation of In-Vitro drug release kinetics of encapsulated '
     'progesterone.'),
    ('Evaluation of encapsulated progesterone based intra-vaginal sponge '
     'delivery system for the management of anestrus and repeat breeding '
     'in cattle.'),
]
for i, txt in enumerate(objectives, start=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p, f'{i}.\t', bold=True)
    add_run(p, txt)


# ===== JUSTIFICATION =====
jh = doc.add_paragraph()
jh.paragraph_format.space_before = Pt(10)
add_run(jh, 'Justification:', bold=True, size=12)

justification_paragraphs = [
    ('Reproductive efficiency is a central determinant of profitability '
     'and sustainability in Indian dairy production systems, where '
     'reproductive disorders alone account for substantial economic losses '
     'every year. Among these, anestrus and repeat breeding syndrome (RBS) '
     'are the two most prevalent reproductive maladies in cyclic and '
     'post-partum cows, lengthening the calving-to-conception interval, '
     'reducing lifetime milk yield and inflating culling rates (Kumaresan '
     'and Srivastava, 2022). Visual estrus detection is particularly '
     'difficult in the Indian sub-tropical environment because peak estrus '
     'expression occurs predominantly during the night and silent estrus '
     'is widespread, leading to mistimed inseminations and a downstream '
     'cascade of reproductive failure (Reith and Hoy, 2018).'),

    ('Progesterone (P4) plays a central regulatory role at the '
     'hypothalamic–pituitary–ovarian (HPO) axis. Sustained luteal-phase '
     'progesterone exposure followed by withdrawal mimics the natural '
     'luteolytic cascade and synchronises follicular waves and ovulation; '
     'this principle, first established by Wiltbank et al. (1961) and '
     'adapted to intravaginal delivery by Roche and Crowley (1973), '
     'continues to underpin all modern oestrus-induction and '
     'oestrus-synchronisation protocols. Macmillan and Peterson (1993) '
     'subsequently operationalised this principle into the Controlled '
     'Internal Drug Release (CIDR) device that, together with the '
     'Progesterone Releasing Intravaginal Device (PRID), is now the global '
     'gold standard for progesterone-based estrus management in cattle '
     'and small ruminants.'),

    ('However, the conventional silicone-based intravaginal progesterone '
     'devices suffer from several well-documented limitations that '
     'compromise both their performance and their sustainability. The '
     'non-biodegradable silicone matrix releases only a fraction of the '
     'loaded hormone and retains a substantial residual load at removal, '
     'generating significant non-biodegradable plastic waste and '
     'hormone-contaminated effluent (Rathbone et al., 2002; Chacher et al., '
     '2017). High device cost and import dependency on Australia, New '
     'Zealand and the United States restrict access in resource-constrained '
     'Indian dairy enterprises. In addition, prolonged silicone retention '
     'may cause vaginal discharge and epithelial irritation, and '
     'inconsistent hormone-release profiles have been reported when devices '
     'are reused across protocols. These collectively underscore the '
     'imperative need for biodegradable, controlled-release progesterone '
     'delivery systems that operate at lower drug doses, lower cost and '
     'reduced environmental burden.'),

    ('Sodium alginate, an anionic algal polysaccharide composed of '
     '(1→4)-linked β-D-mannuronic acid and α-L-guluronic acid residues, is '
     'one of the most attractive candidate biopolymers for such an '
     'application. It is biocompatible, biodegradable, mucoadhesive '
     'through carboxylate–mucin electrostatic and hydrogen-bonding '
     'interactions, and has long enjoyed FDA-GRAS status and an established '
     'pharmaceutical regulatory history (Lai et al., 2024; Colin et al., '
     '2024). Critically, alginate undergoes instantaneous ionotropic '
     'gelation in the presence of divalent cations such as Ca²⁺, forming '
     '"egg-box" hydrogel networks whose density and porosity can be tuned '
     'to modulate drug release (Szekalska et al., 2018a).'),

    ('Spray drying is the formulation method of choice for translating an '
     'aqueous alginate–progesterone feed into a free-flowing dry powder. '
     'It is a rapid, continuous, single-step, scalable process in which '
     'atomised droplets evaporate within milliseconds, producing '
     'reproducible microparticles with controlled size, low residual '
     'moisture and high yield (Broadhead et al., 1992; Vehring, 2008; '
     'Bowey and Neufeld, 2010). The wet-bulb temperature of the drying '
     'droplet remains far below the inlet air temperature, thereby '
     'preserving thermolabile actives. Recent work has demonstrated that '
     'one-step CaCl₂-modified alginate spray drying yields reproducible '
     'mucoadhesive microparticles capable of prolonging drug release '
     '(Szekalska et al., 2018a), and that combining alginate with '
     'hydrophobic-affinity co-polymers such as gelatin (Szekalska et al., '
     '2023), ethyl cellulose (Murata et al., 2009) or hypromellose '
     '(Kruk et al., 2024) enables successful encapsulation of strongly '
     'hydrophobic actives — a critical requirement for a progesterone '
     'formulation, given the steroid\'s poor aqueous solubility '
     '(approximately 7–10 µg/mL) and log P of approximately 3.87.'),

    ('A rigorous understanding of in-vitro release kinetics is a '
     'pre-requisite to successful in-vivo translation. Fitting cumulative '
     'release data to zero-order, first-order, Higuchi, Hixson–Crowell and '
     'Korsmeyer–Peppas models, and computing the diffusion exponent (n), '
     'allows the dominant transport regime to be classified as Fickian '
     'diffusion, anomalous transport, Case-II transport or super Case-II '
     'transport (Peppas et al., 2000). Such mechanistic insight is '
     'essential for engineering a release profile that will support a '
     'near-zero-order plasma progesterone plateau over a 7-day '
     'intravaginal regimen, mirroring the duration of a CIDR insert.'),

    ('Despite the documented success of alginate-based spray-dried '
     'microparticles for several other actives such as metformin, '
     'luliconazole, posaconazole and insulin, and despite the closely '
     'related demonstration by Lopedota et al. (2018) that spray-dried '
     'tripolyphosphate-crosslinked chitosan microparticles can deliver '
     'progesterone for estrus synchronisation in cattle, no peer-reviewed '
     'study has so far reported a sodium-alginate-based, spray-dried, '
     'micronised progesterone microparticulate system incorporated into '
     'an intravaginal sponge specifically for the dual indication of '
     'anestrus and repeat breeding syndrome in dairy cattle. The present '
     'study is therefore designed to fill this gap by integrating '
     'biopolymer encapsulation science, spray-drying process optimisation, '
     'in-vitro release kinetic characterisation and in-vivo cattle '
     'reproductive performance evaluation in a single, translational '
     'research programme tailored to the needs of Indian dairy '
     'production.'),
]

for txt in justification_paragraphs:
    p = justified_para(doc)
    add_run(p, txt)


# ===== SUMMARY OF WORK DONE IN INDIA AND ABROAD =====
sh = doc.add_paragraph()
sh.paragraph_format.space_before = Pt(10)
add_run(sh, 'A summary of work done in India and abroad:', bold=True, size=12)

sh_in = doc.add_paragraph()
add_run(sh_in, 'India', bold=True, italic=True, size=12)

india_paragraphs = [
    ('Indian research on progesterone-mediated estrus management in cattle '
     'and buffalo originated from foundational work in the 1980s when '
     'Rao and Rao (1981) and Sane et al. (1988) reported that exogenous '
     'progesterone treatment could re-initiate cyclicity in anestrous '
     'animals, although delivery was through injectable or oral progestogen '
     'routes that lacked precise control over plasma kinetics.'),

    ('Following the international introduction of intravaginal '
     'progesterone-releasing devices in the 1990s, Indian groups extended '
     'global protocols to tropical field conditions. Kharche et al. (1997) '
     'and Bhadwal et al. (1999) demonstrated acceptable rates of conception '
     'with CIDR/PRID-based estrus induction in cattle and buffalo. '
     'Sah and Nakao (2006) and Patel et al. (2009) further validated '
     'CIDR-based estrus synchronisation regimens in postpartum anestrous '
     'buffalo and dairy cattle, respectively, while documenting practical '
     'drawbacks such as device cost, vaginal discharge and reuse-related '
     'hygiene concerns.'),

    ('The post-2010 Indian literature has focused primarily on '
     'protocol-level optimisation of CIDR-based timed-AI strategies '
     '(Kumar et al., 2013; Kumar and Mandal, 2016) rather than on '
     'innovating the underlying delivery hardware or biopolymer chemistry. '
     'Sridhar et al. (2017) explored conventional polymer implants for '
     'progesterone release, and Ramakrishnan et al. (2019) reported '
     'preliminary work on nanofibre-based progesterone delivery. More '
     'recently, Lavanya et al. (2024) demonstrated that electrospun '
     'pullulan nanofibres (68–123 nm) could encapsulate hydrophobic '
     'progesterone and provide sustained release over seven days '
     'predominantly via Fickian diffusion.'),

    ('However, no published Indian study to date has systematically '
     'optimised sodium-alginate-based, spray-dried, micronised '
     'progesterone microparticles for intravaginal sponge delivery, nor '
     'has any Indian group simultaneously addressed the dual indications '
     'of anestrus and repeat breeding syndrome through such a system. '
     'The present proposal is positioned to bridge this clear research '
     'gap by adapting state-of-the-art alginate spray-drying technology '
     'to the formulation needs of Indian dairy cattle and the specific '
     'reproductive challenges of the sub-tropical production environment.'),
]

for txt in india_paragraphs:
    p = justified_para(doc)
    add_run(p, txt)


sh_ab = doc.add_paragraph()
sh_ab.paragraph_format.space_before = Pt(6)
add_run(sh_ab, 'Abroad', bold=True, italic=True, size=12)

abroad_paragraphs = [
    ('Internationally, the scientific scaffolding for progesterone-mediated '
     'estrus control was laid by Wiltbank et al. (1961), who established '
     'the fundamental principle that exogenous progesterone suppresses '
     'gonadotropin-driven follicular waves and that controlled withdrawal '
     'triggers synchronised ovulation. Roche and Crowley (1973) translated '
     'this principle into intravaginal delivery, and Macmillan and Peterson '
     '(1993) consolidated it into the CIDR device that has since become '
     'the global benchmark for cattle estrus synchronisation.'),

    ('Polymeric controlled-drug-delivery science evolved in parallel. '
     'Langer (1990) systematically established that hydrophobic, slowly '
     'degrading polymeric matrices can sustain bioactive release over '
     'extended periods. Göpferich (1996) characterised the surface- versus '
     'bulk-erosion behaviour of biodegradable polyesters, and Peppas and '
     'co-workers (2000) advanced the diffusion-controlled release theory '
     'that continues to underpin contemporary kinetic modelling. Spray '
     'drying emerged as the scalable, single-step encapsulation method of '
     'choice for pharmaceutical and veterinary formulations following the '
     'work of Broadhead et al. (1992) and the comprehensive '
     'particle-engineering treatment by Vehring (2008). Mora-Huertas et al. '
     '(2010) reinforced the value of high-pressure homogenisation and '
     'ultrasonication as upstream emulsification steps for improving '
     'encapsulation efficiency and reducing particle-size variability.'),

    ('Internationally, sodium alginate has been investigated as a versatile '
     'spray-drying wall material across multiple drug classes. Szekalska '
     'and co-workers reported one-step CaCl₂-modified alginate '
     'microparticles for metformin (2018a) and chitosan-crosslinked '
     'alginate microparticles with prolonged in-vivo metformin release '
     '(Szekalska et al., 2017), and recently demonstrated alginate–gelatin '
     'spray-dried microparticles encapsulating the hydrophobic antifungal '
     'luliconazole (Szekalska et al., 2023). Murata et al. (2009) showed '
     'that alginate combined with ethyl cellulose in a water-in-oil '
     'emulsion spray-drying feed produced pH-responsive microparticles, '
     'and Kruk et al. (2024) used alginate–hypromellose composites for '
     'sustained release of the hydrophobic posaconazole. Meng et al. '
     '(2017) developed thiolated-chitosan-coated alginate multilayer '
     'microparticles by spray drying for vaginal mucoadhesive delivery, '
     'and Lopedota et al. (2018) directly validated spray-dried '
     'tripolyphosphate-crosslinked chitosan microparticles for '
     'progesterone delivery in cattle. Biodegradable intravaginal '
     'progesterone systems in ruminants have additionally been reported '
     'by Searle et al. (2001) and González-Bulnes et al. (2005). Despite '
     'these advances, an integrated alginate spray-dried micronised '
     'progesterone intravaginal sponge specifically validated against '
     'both anestrus and repeat breeding syndrome in cattle remains '
     'unreported in the international literature.'),
]

for txt in abroad_paragraphs:
    p = justified_para(doc)
    add_run(p, txt)


# ===== TECHNICAL PROGRAMME =====
tph = doc.add_paragraph()
tph.paragraph_format.space_before = Pt(10)
add_run(tph, 'Brief Technical programme of work:', bold=True, size=12)

tp_paragraphs = [
    ('Objective 1 – Optimisation of sodium alginate for progesterone '
     'nano/micro-encapsulation by spray drying. Pharmaceutical-grade sodium '
     'alginate (medium G-content) and micronised progesterone (USP/IP '
     'grade) will form the basis of all formulations. Aqueous sodium '
     'alginate solutions (1–3 % w/v) will be prepared and progesterone '
     'introduced either as a homogeneous suspension or via an oil-in-water/'
     'water-in-oil pre-emulsion, with high-speed homogenisation '
     '(10,000–20,000 rpm, 5 min) followed by probe ultrasonication '
     '(40 % amplitude, 1–3 min) to refine droplet size. Where indicated, '
     'in-situ ionotropic crosslinking will be achieved by addition of '
     'CaCl₂ (0.05–0.10 % w/v) to the feed. Spray drying will be performed '
     'on a Büchi B-290 mini spray dryer (or equivalent), with the '
     'following systematically varied parameters: inlet air temperature '
     '(120–160 °C), feed flow rate (3–6 mL/min), atomising air flow '
     '(500–700 NL/h) and aspirator setting (90–100 %). A factorial / '
     'Box–Behnken design-of-experiments approach will identify the '
     'optimal operating window. Physicochemical characterisation will '
     'include particle-size distribution and polydispersity index (PDI) '
     'by laser diffraction, zeta potential, surface morphology by '
     'scanning electron microscopy (SEM), encapsulation efficiency (EE %) '
     'and drug loading (DL %) by HPLC after solvent extraction, thermal '
     'behaviour by differential scanning calorimetry (DSC), crystallinity '
     'by X-ray diffraction (XRD) and drug–polymer compatibility by '
     'Fourier-transform infrared spectroscopy (FTIR).'),

    ('Objective 2 – In-vitro drug release kinetics. The optimised '
     'formulation will be subjected to dissolution studies in simulated '
     'vaginal fluid (SVF, pH ~4.5) and phosphate-buffered saline (PBS, '
     'pH 7.4) at 37 °C using a USP dissolution apparatus or a dialysis-bag '
     'set-up. Aliquots will be withdrawn at 0, 1, 2, 4, 8, 12, 24, 48, 72 '
     'and 168 h, and progesterone quantified by validated UV-Vis or HPLC '
     'methods. Cumulative release data will be fitted to zero-order, '
     'first-order, Higuchi, Hixson–Crowell and Korsmeyer–Peppas equations '
     'to identify the dominant release mechanism, and the diffusion '
     'exponent (n) will be computed to discriminate between Fickian, '
     'anomalous, Case-II and super Case-II transport regimes. Cytotoxicity '
     'and biocompatibility of the optimised formulation will be assessed '
     'on a relevant epithelial cell line by MTT and LDH assays.'),

    ('Objective 3 – In-vivo evaluation of intravaginal sponge in anestrus '
     'and repeat breeding cattle. The optimised progesterone-loaded sodium '
     'alginate microparticles will be incorporated into polyurethane-based '
     'intravaginal sponges. Three groups of cattle will be enrolled: '
     '(i) clinically diagnosed anestrous cows (n = 12), (ii) repeat '
     'breeder cows (n = 12) and (iii) a positive-control group on a '
     'commercial CIDR insert (n = 12). Outcome variables will include '
     'estrus response rate (%), onset of estrus post-device removal, '
     'duration of estrus, plasma progesterone profile by RIA / ELISA at '
     'predefined intervals, ovarian follicular dynamics by transrectal '
     'ultrasonography, conception rate at first artificial insemination '
     'and repeat-breeding correction rate at 60-day pregnancy diagnosis. '
     'Statistical analysis will employ Chi-square / Fisher\'s exact tests '
     'for categorical proportions and one-way ANOVA with Tukey\'s '
     'post-hoc test for continuous variables, with statistical '
     'significance set at P < 0.05.'),
]

for txt in tp_paragraphs:
    p = justified_para(doc)
    add_run(p, txt)


# ===== LOCATION / FACILITIES / COLLABORATION =====
labelled_field(doc, 'Location and place of work:',
               'Dairy Production Section, Southern Regional Station, '
               'ICAR-National Dairy Research Institute, Bengaluru, '
               'Karnataka.')
labelled_field(doc, 'Facilities available:',
               'All the facilities required for this research work are '
               'available at SRS-ICAR-NDRI, Bengaluru, Karnataka.')
labelled_field(doc, 'Collaboration with other Divisions/Institutes:',
               'The proposed research work seeks collaboration from the '
               'Animal Biochemistry and Dairy Engineering Divisions of '
               'this institute.')


# ===== ADVISORY COMMITTEE TABLE =====
acp = doc.add_paragraph()
acp.paragraph_format.space_before = Pt(12)
acp.paragraph_format.space_after = Pt(4)
add_run(acp, 'Advisory Committee', bold=True, size=12)

committee = [
    ('Dr. Jeyakumar Sakthivel (Major Advisor)', 'Major',
     'Principal Scientist'),
    ('Dr. A. Kumaresan (Member)', 'Major',
     'ICAR-National Fellow and Principal Scientist'),
    ('Dr. Rubina Kumari Baithalu (Member)', 'Major',
     'Senior Scientist'),
    ('Dr. Vedamurthy, G. V. (Member)', 'Allied',
     'Senior Scientist'),
    ('Dr. F. Magdaline Eljeeva Emerald (Member)', 'Allied',
     'Principal Scientist'),
]

table = doc.add_table(rows=1 + len(committee) + 1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

# header row
header_cells = table.rows[0].cells
headers = ['Name of committee member', 'Field\n(Major/Allied)',
           'Designation', 'Signature with date']
for j, h in enumerate(headers):
    cell = header_cells[j]
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, h, bold=True, size=11)
    set_cell_borders(cell)
    set_cell_shading(cell, 'D9E2F3')
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# committee rows
for i, (name, field, desig) in enumerate(committee, start=1):
    row = table.rows[i].cells
    for j, val in enumerate([name, field, desig, '']):
        cell = row[j]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 \
            else WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, val, size=11)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# director's nominee row
dn_row = table.rows[1 + len(committee)].cells
for j, val in enumerate(["Director's Nominee", '', '', '']):
    cell = dn_row[j]
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 \
        else WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, val, bold=(j == 0), size=11)
    set_cell_borders(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# column widths
col_widths = [Cm(6.5), Cm(2.8), Cm(4.5), Cm(3.5)]
for j, w in enumerate(col_widths):
    for row in table.rows:
        row.cells[j].width = w


# ===== SIGNATURE OF HEAD OF DIVISION =====
for _ in range(4):
    doc.add_paragraph()

sig = doc.add_paragraph()
sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_run(sig, 'Signature of Head of Division with Seal',
        bold=True, size=12)


# ===== SAVE =====
out_path = ('/projects/sandbox/MAIN-KIRO/'
            'Format_Approval_Title_Advisory_Committee_Alginate_P4.docx')
doc.save(out_path)
print(f'Saved: {out_path}')
